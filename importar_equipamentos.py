import sys
import unicodedata
from collections import Counter, defaultdict

from openpyxl import load_workbook
from psycopg2.extras import execute_values

from db import get_connection
from status import normalizar_status

# ============================================================
# CONFIGURAÇÕES — ajuste se necessário
# ============================================================
EXCEL_PATH = "RELACAO_INVENTARIO_EDITADO.xlsx"

# A planilha quase não tem STATUS preenchido. Linhas com STATUS em branco
# entram com o status abaixo (troque para "FUNCIONAL" se preferir).
STATUS_PADRAO = "COMPLETO"
# ============================================================


def so_letras(texto):
    """'SÉRIE' -> 'SERIE', 'Nº' -> 'NO', 'OBSERVAÇÃO' -> 'OBSERVACAO'."""
    sem_acento = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in sem_acento if c.isalpha() and ord(c) < 128).upper()


def texto_ou_none(val):
    """Converte vazio/None em None e tira espaços das pontas."""
    if val is None:
        return None
    val = str(val).strip()
    return val or None


def achar_cabecalho(ws):
    """Procura a linha do cabeçalho (a que tem SÉRIE) e devolve (linha, {nome: coluna})."""
    for r in range(1, 15):
        colunas = {}
        for c in range(1, ws.max_column + 1):
            nome = so_letras(ws.cell(r, c).value or "")
            if nome:
                colunas[nome] = c
        if "SERIE" in colunas and "EQUIPAMENTO" in colunas:
            return r, colunas
    return None, None


def ler_aba(ws):
    """Lê uma aba e devolve a lista de registros (dicts)."""
    linha_cab, cols = achar_cabecalho(ws)
    if linha_cab is None:
        print(f"   [aviso] aba '{ws.title}' ignorada: cabeçalho SÉRIE/EQUIPAMENTO não encontrado.")
        return []

    col_num = next((cols[k] for k in ("NO", "N", "NUMERO", "NRO") if k in cols), None)
    col_status = cols.get("STATUS")
    col_obs = cols.get("OBSERVACAO")

    # Aba "N - Carrinho 1" -> tipo NOTEBOOK, carrinho "Carrinho 1"
    # Aba "Tablet - Acessa" -> tipo TABLET,   carrinho "Acessa"
    titulo = f"{ws['A1'].value or ''} {ws.title}".upper()
    tipo = "TABLET" if "TABLET" in titulo else "NOTEBOOK"
    carrinho = ws.title.split(" - ", 1)[1].strip() if " - " in ws.title else ws.title.strip()

    registros = []
    numero_anterior = 0
    for r in range(linha_cab + 1, ws.max_row + 1):
        serie = texto_ou_none(ws.cell(r, cols["SERIE"]).value)
        equip = texto_ou_none(ws.cell(r, cols["EQUIPAMENTO"]).value)
        status_txt = texto_ou_none(ws.cell(r, col_status).value) if col_status else None
        obs = texto_ou_none(ws.cell(r, col_obs).value) if col_obs else None

        if not (serie or equip or status_txt or obs):
            continue  # linha em branco

        # Nº: na planilha é uma fórmula (=SUM(A4+1)); usa o valor se for número,
        # senão continua a sequência do carrinho.
        valor_num = ws.cell(r, col_num).value if col_num else None
        if isinstance(valor_num, (int, float)):
            numero = int(valor_num)
        else:
            numero = numero_anterior + 1
        numero_anterior = numero

        if serie in ("-", "—"):
            serie = None
        if serie:
            serie = serie.upper()

        if status_txt:
            status = normalizar_status(status_txt)
            if status is None:
                raise ValueError(
                    f"Status desconhecido '{status_txt}' na aba '{ws.title}', linha {r}. "
                    "Corrija na planilha (completo, funcional, em análise, reparo, "
                    "não funcional ou desaparecido)."
                )
        else:
            status = STATUS_PADRAO

        registros.append({
            "tipo": tipo, "carrinho": carrinho, "numero": numero,
            "serie": serie, "equipamento": equip or "NÃO INFORMADO",
            "status": status, "observacao": obs,
            "status_em_branco": status_txt is None,
            "linha": r, "aba": ws.title,
        })
    return registros


def importar(substituir=False):
    # 1. Lê o Excel
    print("Lendo o arquivo Excel...")
    wb = load_workbook(EXCEL_PATH)
    registros = []
    for ws in wb.worksheets:
        lidos = ler_aba(ws)
        print(f"   {ws.title}: {len(lidos)} equipamentos")
        registros.extend(lidos)
    print(f"   Total: {len(registros)} equipamentos.\n")

    # 2. Relatório de conferência dos dados
    print("Conferência dos dados:")
    por_status = Counter(r["status"] for r in registros)
    em_branco = sum(1 for r in registros if r["status_em_branco"])
    print(f"   Status: {dict(por_status)}")
    print(f"   {em_branco} linhas sem STATUS na planilha entraram como {STATUS_PADRAO}.")

    sem_serie = [r for r in registros if not r["serie"]]
    print(f"   {len(sem_serie)} sem SÉRIE: " +
          "; ".join(f"{r['aba']} linha {r['linha']}" for r in sem_serie))

    locais = defaultdict(list)
    for r in registros:
        if r["serie"]:
            locais[r["serie"]].append(f"{r['aba']} linha {r['linha']}")
    duplicadas = {s: l for s, l in locais.items() if len(l) > 1}
    for s, l in duplicadas.items():
        print(f"   [confira] SÉRIE {s} repetida em: " + "; ".join(l))
    print()

    # 3. Conecta no banco
    print("Conectando ao PostgreSQL...")
    conn = get_connection()
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM equipamentos;")
    existentes = cur.fetchone()[0]
    if existentes > 0 and not substituir:
        cur.close()
        conn.close()
        print(f"\nA tabela já tem {existentes} equipamentos (com possíveis alterações feitas no sistema).")
        print("Nada foi alterado. Para apagar e recarregar da planilha, rode:")
        print("   python importar_equipamentos.py --substituir")
        sys.exit(1)

    # 4. Backup antes de apagar
    print("Criando backup da tabela atual (equipamentos_backup)...")
    cur.execute("DROP TABLE IF EXISTS equipamentos_backup;")
    cur.execute("CREATE TABLE equipamentos_backup AS SELECT * FROM equipamentos;")
    print("   Backup criado com sucesso.\n")

    # 5. Apaga registros antigos e reinicia o ID
    print("Removendo registros antigos...")
    cur.execute("TRUNCATE TABLE equipamentos RESTART IDENTITY;")

    # 6. Insere os novos registros
    print("Inserindo novos registros...")
    rows = [(r["tipo"], r["carrinho"], r["numero"], r["serie"], r["equipamento"],
             r["status"], r["observacao"]) for r in registros]
    execute_values(cur, """
        INSERT INTO equipamentos
            (tipo, carrinho, numero, serie, equipamento, status, observacao)
        VALUES %s
    """, rows)

    conn.commit()
    cur.close()
    conn.close()

    print(f"\nPronto! {len(rows)} equipamentos importados com sucesso.")
    print("   Acesse http://localhost:5000 para conferir.")


if __name__ == "__main__":
    args = sys.argv[1:]
    substituir = "--substituir" in args
    caminhos = [a for a in args if not a.startswith("--")]
    if caminhos:
        EXCEL_PATH = caminhos[0]
    importar(substituir)
